# -*- coding: utf-8 -*-
"""
Word (.docx) 文件文本与公式提取器 (Docx & Math Extractor)

支持提取普通文本段落、表格单元格以及 Word 原生数学公式 (OMML / m:oMath)。
"""

import logging
from typing import List
from docx import Document

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_path: str) -> str:
    """从 docx 文件中提取完整文本，包括段落中的普通文本和 Word OMML 公式文本。

    相比默认的 paragraph.text（仅提取 w:t），本函数会遍历 XML 树中的所有文本节点
    （包含 w:t 和 m:t 等），确保数学公式和推导符号不会被遗漏。
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        logger.error(f"读取 docx 文件失败 ({file_path}): {e}")
        raise

    extracted_lines: List[str] = []

    # 1. 提取所有段落内容
    for p in doc.paragraphs:
        parts: List[str] = []
        for elem in p._element.iter():
            tag = elem.tag
            if tag.endswith("}t") and elem.text:
                parts.append(elem.text)
            elif tag.endswith("}tab"):
                parts.append("\t")
            elif tag.endswith("}br") or tag.endswith("}cr"):
                parts.append("\n")
        line = "".join(parts)
        extracted_lines.append(line)

    # 2. 如果包含表格，提取表格内容（避免丢失表格中的文本或公式）
    if doc.tables:
        for table in doc.tables:
            for row in table.rows:
                row_cells_text = []
                for cell in row.cells:
                    cell_parts = []
                    for elem in cell._element.iter():
                        tag = elem.tag
                        if tag.endswith("}t") and elem.text:
                            cell_parts.append(elem.text)
                        elif tag.endswith("}tab"):
                            cell_parts.append("\t")
                        elif tag.endswith("}br") or tag.endswith("}cr"):
                            cell_parts.append("\n")
                    row_cells_text.append("".join(cell_parts).strip())
                extracted_lines.append("\t".join(row_cells_text))

    return "\n".join(extracted_lines)
