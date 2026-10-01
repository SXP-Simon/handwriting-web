# -*- coding: utf-8 -*-
import os
import pytest
from docx import Document
from backend.docx_extractor import extract_text_from_docx


def test_extract_text_from_docx_basic(tmp_path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("第一行测试文本")
    doc.add_paragraph("第二行测试文本\t带制表符")
    doc.save(str(doc_path))

    extracted = extract_text_from_docx(str(doc_path))
    assert "第一行测试文本" in extracted
    assert "第二行测试文本\t带制表符" in extracted


def test_extract_text_from_exec_docx_if_present():
    exec_docx_path = r"C:\Users\Simon\Downloads\exec.docx"
    if not os.path.exists(exec_docx_path):
        pytest.skip("exec.docx not found at local downloads path")

    extracted = extract_text_from_docx(exec_docx_path)
    # 验证关键数学公式与推导符号完整提取
    assert "LG6={a∣a∈Σ" in extracted or "LG6={a" in extracted
    assert "N⇒ND⇒NDD⇒NDDD" in extracted
    assert "S→1∣3∣5∣7∣9" in extracted
    assert "E⇒E+T⇒T+T" in extracted
